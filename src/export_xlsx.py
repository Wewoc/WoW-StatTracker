#!/usr/bin/env python3
# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (C) 2026 Wewoc (github.com/wewoc)
# export_xlsx.py
# Excel export: reads all wow_stats*.db (in data/) and writes stattracker.xlsx.
#
# One sheet per game version (wow_stats.db -> "Retail", wow_stats_anniversary.db -> "Anniversary" ...).
# Per sheet: one row per statistic, one column per character (newest scan each).
# Statistics that are empty for all characters ('--') are left out.
#
# Units (internal codes -> translated labels in the "Unit" column):
#   num   = raw value
#   gold  = copper / 10000 (gold as a decimal number)
#   hours = seconds / 3600 (played time only)
#   text  = not numeric (e.g. '5164 (Humanoid)')
#   mix   = row contains numbers AND text for different characters
#
# Needs: pip install openpyxl (bundled in the exe release)
# Usage: python export_xlsx.py [--output PATH]

import argparse
import re
import sqlite3
import sys
from pathlib import Path

import paths
import stat_parser
from i18n import t

GOLD_RE = re.compile(r"\d[\d.,]*\s*[gsc]\b")
PLAYED_IDS = tuple(sid for sid, _f, _k in stat_parser.PLAYED_STATS)
EMPTY = ("--", "")


def default_out():
    return paths.data_dir() / "stattracker.xlsx"


def sheet_title(db_path):
    stem = db_path.stem  # wow_stats / wow_stats_anniversary
    suffix = stem[len("wow_stats"):].strip("_")
    return "Retail" if not suffix else suffix.replace("_", " ").title()


def find_dbs(folder=None):
    folder = Path(folder) if folder else paths.data_dir()
    return sorted(folder.glob("wow_stats*.db"), key=lambda p: (p.stem != "wow_stats", p.stem))


def load_db(db_path):
    """Returns (chars, stats): chars = list of dicts, stats = list (id, name, cat, {key: (text, num)})."""
    conn = sqlite3.connect(db_path)
    try:
        stat_parser.check_schema(conn, db_path)      # a database from a newer program is not read
        scans = conn.execute("""
            SELECT s.scan_id, s.char_key, s.scanned_at, s.class, s.level
            FROM scans s
            WHERE s.scan_id = (SELECT scan_id FROM scans WHERE char_key = s.char_key
                               ORDER BY scanned_at DESC, scan_id DESC LIMIT 1)
            ORDER BY s.char_key""").fetchall()
        chars = [{"scan_id": r[0], "key": r[1], "scanned_at": r[2], "class": r[3], "level": r[4]} for r in scans]
        by_scan = {c["scan_id"]: c["key"] for c in chars}
        defs = {r[0]: [r[0], r[1], r[2], {}] for r in conn.execute("SELECT stat_id,name,category FROM stat_defs")}
        for sid, (name, cat) in stat_parser.played_labels().items():
            if sid in defs:
                defs[sid][1], defs[sid][2] = name, cat
        for scan_id, stat_id, text, num in conn.execute(
                "SELECT scan_id,stat_id,value_text,value_num FROM stat_values"):
            if scan_id in by_scan and stat_id in defs:
                defs[stat_id][3][by_scan[scan_id]] = (text, num)
    finally:
        conn.close()
    return chars, sorted(defs.values(), key=lambda d: (d[2], d[1]))


def build_rows(chars, stats):
    """Turns statistics into rows: (cat, name, unit_code, differs, [cell values], best_or_None).
    Cell value = number (converted), text or None."""
    rows = []
    for stat_id, name, cat, vals in stats:
        cells = [vals.get(c["key"]) for c in chars]
        if all((x is None or x[0] in EMPTY) for x in cells):
            continue

        texts = {(x[0] if x else "--") for x in cells}
        differs = len(texts) > 1

        has_gold = any(x and GOLD_RE.search(x[0]) for x in cells)
        if stat_id in PLAYED_IDS:
            unit, conv = "hours", lambda n: n / 3600
        elif has_gold:
            unit, conv = "gold", lambda n: n / 10000
        else:
            unit, conv = "num", lambda n: n

        out, kinds = [], set()
        for x in cells:
            if x is None or x[0] in EMPTY:
                out.append(None)
            elif x[1] is not None:
                out.append(conv(x[1]))
                kinds.add("num")
            else:
                out.append(x[0])
                kinds.add("text")
        if kinds == {"text"}:
            unit = "text"
        elif kinds == {"num", "text"}:
            unit = "mix"

        best = None
        nums = [v for v in out if isinstance(v, (int, float))]
        if len(chars) > 1 and len(nums) == len(out) and len(set(nums)) > 1:
            best = max(nums)
        rows.append((cat, name, unit, differs, out, best))
    return rows


def export(out_path=None, folder=None):
    """Writes the workbook. Returns a list of info lines. Raises SystemExit(str) on errors."""
    out_path = Path(out_path) if out_path else default_out()
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit(t("xl_err_openpyxl"))

    folder = Path(folder) if folder else paths.data_dir()
    dbs = find_dbs(folder)
    if not dbs:
        sys.exit(t("xl_err_no_db", folder=folder))

    fmt = {"num": "#,##0", "gold": "#,##0.00", "hours": "#,##0.0"}
    head_fill = PatternFill("solid", fgColor="1E1E1E")
    head_font = Font(bold=True, color="F4C430")
    best_fill = PatternFill("solid", fgColor="C6EFCE")

    wb = Workbook()
    wb.remove(wb.active)
    info = []
    for db in dbs:
        chars, stats = load_db(db)
        if not chars:
            info.append(t("xl_skip_no_scans", name=db.name))
            continue
        rows = build_rows(chars, stats)
        ws = wb.create_sheet(sheet_title(db)[:31])

        header = [t("xl_col_category"), t("xl_col_stat"), t("xl_col_unit"), t("xl_col_diff")]
        for c in chars:
            parts = [c["key"], " ".join(str(p) for p in (c["class"], c["level"]) if p), c["scanned_at"]]
            header.append("\n".join(p for p in parts if p))
        ws.append(header)
        for cell in ws[1]:
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[1].height = 48

        for cat, name, unit, differs, vals, best in rows:
            ws.append([cat, name, t(f"unit_{unit}"), t("xl_yes") if differs else t("xl_no")] + vals)
            r = ws.max_row
            has_frac = unit == "num" and any(isinstance(v, float) and not v.is_integer() for v in vals)
            for j, v in enumerate(vals):
                cell = ws.cell(row=r, column=5 + j)
                if isinstance(v, (int, float)):
                    cell.number_format = "#,##0.00" if has_frac else fmt[unit if unit in fmt else "num"]
                    if best is not None and v == best:
                        cell.fill = best_fill
                        cell.font = Font(bold=True)

        widths = {"A": 30, "B": 52, "C": 9, "D": 12}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        for j in range(len(chars)):
            ws.column_dimensions[get_column_letter(5 + j)].width = 20
        ws.freeze_panes = "E2"
        ws.auto_filter.ref = f"A1:{get_column_letter(4 + len(chars))}{max(ws.max_row, 2)}"
        info.append(t("xl_sheet_info", sheet=ws.title, chars=len(chars), stats=len(rows), db=db.name))

    if not wb.sheetnames:
        sys.exit(t("xl_err_no_scans"))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(out_path)
    except PermissionError:
        sys.exit(t("xl_err_permission", path=out_path))
    info.append(t("xl_file", path=out_path))
    return info


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--output", default=str(default_out()))
    a = ap.parse_args()
    for line in export(Path(a.output)):
        print(line)


if __name__ == "__main__":
    main()
